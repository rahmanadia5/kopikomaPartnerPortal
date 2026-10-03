"""Build the PRD .docx (Arial 11, justified, header top-left) from PRD.md."""
import re, sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

NAME, UNIV = 'Nadia Rahma Prasanti', 'Universitas Indonesia'
HERE = Path(__file__).parent
OUT = HERE / f'EDTS APM 2027 - {NAME} - {UNIV}.docx'
FONT = 'Arial'

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.2)
    s.left_margin = s.right_margin = Cm(2.2)

def set_font(style_or_run, size=11, bold=None):
    f = style_or_run.font
    f.name, f.size = FONT, Pt(size)
    if bold is not None: f.bold = bold
    f.color.rgb = RGBColor(0, 0, 0)
    el = style_or_run.element if hasattr(style_or_run, 'element') else style_or_run._element
    rpr = el.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rfonts.set(qn(k), FONT)

set_font(doc.styles['Normal'])
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
for name, size in (('Title', 16), ('Heading 1', 14), ('Heading 2', 12), ('Heading 3', 11), ('List Bullet', 11), ('List Bullet 2', 11), ('List Bullet 3', 11), ('List Number', 11)):
    set_font(doc.styles[name], size, bold=name.startswith(('Title', 'Heading')))

hp = doc.sections[0].header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(hp.add_run(f'{NAME} - {UNIV}'))

def add_inline(p, text, size=11):
    for part in re.split(r'(\*\*[^*]+\*\*)', text):
        if not part: continue
        bold = part.startswith('**')
        r = p.add_run(part[2:-2] if bold else part)
        set_font(r, size, bold=bold)

def para(text, style=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    add_inline(p, text)
    return p

def shade(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hex_fill)
    tcpr.append(sh)

def table(rows):
    rows = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
    rows = [r for r in rows if not all(re.fullmatch(r':?-+:?', c) for c in r)]
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    for i, r in enumerate(rows):
        for j, c in enumerate(r):
            cell = t.cell(i, j); cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c[:1] in ('✅', '⚠', '❌') or (i == 0 and rows[0][0] == 'Capability' and j in (1, 2)) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, f'**{c}**' if i == 0 and c else c)
            if i == 0: shade(cell, 'E7E2DA')
    widths = {'Capability': (4.8, 2.9, 1.7, 7.2), 'ID': (2.3, 4.4, 5.0, 4.9), '#': None}.get(rows[0][0])
    if rows[0][0] == '#' and len(rows[0]) == 5: widths = (1.0, 3.6, 4.0, 1.8, 6.0)
    if widths:
        t.autofit = False
        for row in t.rows:
            for j, w in enumerate(widths): row.cells[j].width = Cm(w)
        for j, gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'), str(int(widths[j] * 567)))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

lines = (HERE / 'PRD.md').read_text(encoding='utf-8').splitlines()
toc = [l for l in lines if l.startswith(('## ', '### ', '#### ')) and 'Table of Contents' not in l]

def grid(spec):
    items = [x.strip().split('|') for x in spec.split(';;')]
    t = doc.add_table(rows=2, cols=len(items))
    width = 16.4 / len(items) - 0.4
    for j, (path, cap) in enumerate(items):
        c = t.cell(0, j); c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].add_run().add_picture(str(HERE / path.strip()), width=Cm(width))
        q = t.cell(1, j).paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = q.add_run(cap.strip()); set_font(r, 10); r.italic = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith('|'):
        block = []
        while i < len(lines) and lines[i].startswith('|'): block.append(lines[i]); i += 1
        table(block); continue
    m = re.match(r'!\[[^\]]*\]\(([^)]+)\)', ln)
    if m:
        doc.add_picture(str(HERE / m.group(1)), width=Cm(14.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1; continue
    if ln.strip() == '[[TOC]]':
        for h in toc:
            level = len(h) - len(h.lstrip('#')) - 1
            p = doc.add_paragraph(style=['List Bullet', 'List Bullet 2', 'List Bullet 3'][level - 1])
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, h.lstrip('# '))
        i += 1; continue
    if ln.startswith('[[GRID '):
        grid(ln[7:-2]); i += 1; continue
    if ln.strip() == '[[PAGEBREAK]]':
        doc.add_page_break(); i += 1; continue
    if ln.startswith('# '): p = doc.add_paragraph(style='Title'); add_inline(p, ln[2:], 16)
    elif ln.startswith('## '): doc.add_heading(ln[3:], level=1)
    elif ln.startswith('### '): doc.add_heading(ln[4:], level=2)
    elif ln.startswith('#### '): doc.add_heading(ln[5:], level=3)
    elif ln.startswith('- '): para(ln[2:], 'List Bullet')
    elif re.match(r'\d+\. ', ln): para(re.sub(r'^\d+\. ', '', ln), 'List Number', WD_ALIGN_PARAGRAPH.LEFT)
    elif ln.strip(): para(ln, align=WD_ALIGN_PARAGRAPH.LEFT if ln.startswith(('Prototype', 'Demo password')) else WD_ALIGN_PARAGRAPH.JUSTIFY)
    i += 1

for p in doc.paragraphs:
    if p.style.name.startswith('Heading'):
        for r in p.runs: set_font(r, p.style.font.size.pt, bold=True)

doc.save(OUT)
print(OUT)
